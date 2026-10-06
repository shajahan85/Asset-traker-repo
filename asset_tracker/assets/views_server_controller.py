"""
Server Controller Management Views
Handles 192.168.50.x IP range systems
"""

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.views.generic import ListView, UpdateView, View
from django.urls import reverse_lazy
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from .models import Asset, HyperVVM
from .permissions import AdminRequiredMixin
from .services.asset_service import AssetService
from .services.export_service import ExportService
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO


class ServerControllerListView(AdminRequiredMixin, ListView):
    """
    Display Server Controllers (192.168.50.x systems).
    Shows both active and freed in tabbed interface.
    """
    model = Asset
    template_name = 'assets/server_controller_list.html'
    context_object_name = 'server_controllers'
    
    def get_queryset(self):
        """Get server controllers based on active/freed tab and IP filter."""
        # Get tab parameter (default: active)
        tab = self.request.GET.get('tab', 'active')
        ip_filter = self.request.GET.get('ip_filter', '').strip()
        
        # Base queryset: all assets
        queryset = Asset.objects.select_related(
            'operating_system', 'team', 'branch'
        ).prefetch_related('asset_ip_addresses__ip_address', 'hosted_vms')
        
        # If IP filter is provided, handle VM IP filtering
        if ip_filter:
            # Check if this IP belongs to a VM
            vm_with_ip = HyperVVM.objects.filter(
                Q(ip_address__address__icontains=ip_filter) |
                Q(manual_ip__icontains=ip_filter)
            ).first()
            
            if vm_with_ip:
                # If IP belongs to a VM, show the host machine
                queryset = queryset.filter(pk=vm_with_ip.host_machine.pk)
            else:
                # Otherwise filter by asset IP
                queryset = queryset.filter(
                    Q(asset_ip_addresses__ip_address__address__icontains=ip_filter) |
                    Q(manual_ip__icontains=ip_filter)
                ).distinct()
        else:
            # No IP filter: show all 192.168.50.x systems
            queryset = queryset.filter(
                Q(asset_ip_addresses__ip_address__address__startswith='192.168.50.') |
                Q(manual_ip__startswith='192.168.50.')
            ).distinct()
        
        # Filter by status based on tab
        if tab == 'freed':
            queryset = queryset.filter(status='freed')
        else:
            queryset = queryset.filter(status='active')
        
        # Apply search if provided
        search = self.request.GET.get('search', '').strip()
        if search:
            queryset = queryset.filter(
                Q(asset_tag__icontains=search) |
                Q(assigned_to__icontains=search)
            )
        
        # Apply team filter if provided
        team_id = self.request.GET.get('team')
        if team_id:
            queryset = queryset.filter(team_id=team_id)
        
        # Apply branch filter if provided
        branch_id = self.request.GET.get('branch')
        if branch_id:
            queryset = queryset.filter(branch_id=branch_id)
        
        return queryset.order_by('asset_tag')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['current_tab'] = self.request.GET.get('tab', 'active')
        context['ip_filter'] = self.request.GET.get('ip_filter', '')
        
        # Check if filtering by VM IP
        ip_filter = self.request.GET.get('ip_filter', '').strip()
        if ip_filter:
            vm_with_ip = HyperVVM.objects.filter(
                Q(ip_address__address__icontains=ip_filter) |
                Q(manual_ip__icontains=ip_filter)
            ).first()
            if vm_with_ip:
                context['filtered_by_vm'] = True
                context['filtered_vm_ip'] = ip_filter
                context['filtered_vm_name'] = vm_with_ip.name
        
        # Get filter options
        from .models import Team, Branch
        context['teams'] = Team.objects.all().order_by('name')
        context['branches'] = Branch.objects.all().order_by('name')
        
        # Get counts for tabs
        active_count = Asset.objects.filter(
            Q(asset_ip_addresses__ip_address__address__startswith='192.168.50.') |
            Q(manual_ip__startswith='192.168.50.'),
            status='active'
        ).distinct().count()
        
        freed_count = Asset.objects.filter(
            Q(asset_ip_addresses__ip_address__address__startswith='192.168.50.') |
            Q(manual_ip__startswith='192.168.50.'),
            status='freed'
        ).distinct().count()
        
        context['active_count'] = active_count
        context['freed_count'] = freed_count
        context['search_query'] = self.request.GET.get('search', '')
        context['selected_team'] = self.request.GET.get('team', '')
        context['selected_branch'] = self.request.GET.get('branch', '')
        
        return context


class ServerControllerEditView(AdminRequiredMixin, UpdateView):
    """Edit Server Controller details."""
    model = Asset
    template_name = 'assets/server_controller_form.html'
    fields = [
        'asset_tag', 'system_type', 'hardware_serial_number',
        'operating_system', 'team', 'branch', 'particulars',
        'warranty_expiration', 'manufacturer'
    ]
    success_url = reverse_lazy('server_controller_list')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['hyperv_vms'] = self.object.hosted_vms.all()
        context['is_server_controller'] = self.object.is_server_controller()
        return context
    
    def form_valid(self, form):
        # Check if IP is changing out of 50.x range
        old_ip = self.object.get_ip_address_string()
        
        response = super().form_valid(form)
        
        # Check new IP after save
        new_ip = self.object.get_ip_address_string()
        
        # Warn if moved out of range
        if old_ip.startswith('192.168.50.') and not new_ip.startswith('192.168.50.'):
            messages.warning(
                self.request,
                f'This system has been removed from Server Controllers list '
                f'as its IP ({new_ip}) is no longer in the 192.168.50.x range.'
            )
        
        messages.success(self.request, f'Server Controller "{self.object.asset_tag}" updated successfully.')
        return response


class ServerControllerFreeView(AdminRequiredMixin, View):
    """Mark a Server Controller as freed."""
    
    def post(self, request, pk):
        asset = get_object_or_404(Asset, pk=pk)
        
        # Verify it's a server controller
        if not asset.is_server_controller():
            messages.error(request, 'This asset is not a Server Controller.')
            return redirect('asset_list')
        
        # Free the asset
        asset.status = 'freed'
        asset.freed_date = timezone.now()
        asset.assigned_to = None
        asset.employee_code = None
        asset.team = None
        asset.save()
        
        # Log history
        from .services.asset_service import AssetService
        AssetService.log_history(
            asset=asset,
            action='freed',
            user=request.user,
            details=f'Server Controller freed: {asset.asset_tag}'
        )
        
        messages.success(
            request,
            f'Server Controller "{asset.asset_tag}" has been freed successfully.'
        )
        return redirect('server_controller_list' + '?tab=freed')


class ServerControllerReassignView(AdminRequiredMixin, View):
    """Reassign a freed Server Controller."""
    
    def get(self, request, pk):
        asset = get_object_or_404(Asset, pk=pk, status='freed')
        
        if not asset.is_server_controller():
            messages.error(request, 'This asset is not a Server Controller.')
            return redirect('freed_systems')
        
        from .models import Team, Branch
        context = {
            'asset': asset,
            'teams': Team.objects.all().order_by('name'),
            'branches': Branch.objects.all().order_by('name'),
        }
        return render(request, 'assets/server_controller_reassign.html', context)
    
    def post(self, request, pk):
        asset = get_object_or_404(Asset, pk=pk, status='freed')
        
        if not asset.is_server_controller():
            messages.error(request, 'This asset is not a Server Controller.')
            return redirect('freed_systems')
        
        # Get form data
        assigned_to = request.POST.get('assigned_to', '').strip()
        team_id = request.POST.get('team')
        notes = request.POST.get('notes', '').strip()
        
        if not assigned_to or not team_id:
            messages.error(request, 'Assigned To and Team are required.')
            return redirect('server_controller_reassign', pk=pk)
        
        # Update asset
        asset.status = 'active'
        asset.assigned_to = assigned_to
        asset.team_id = team_id
        asset.freed_date = None
        asset.health_status = None
        asset.issues_description = None
        
        if notes:
            asset.particulars = notes
        
        asset.save()
        
        # Log history
        from .services.asset_service import AssetService
        AssetService.log_history(
            asset=asset,
            action='reassigned',
            user=request.user,
            details=f'Server Controller reassigned to {assigned_to}'
        )
        
        messages.success(
            request,
            f'Server Controller "{asset.asset_tag}" has been reassigned to {assigned_to}.'
        )
        return redirect('server_controller_list')


class ServerControllerScrapView(AdminRequiredMixin, View):
    """Scrap a freed Server Controller."""
    
    def post(self, request, pk):
        asset = get_object_or_404(Asset, pk=pk, status='freed')
        
        if not asset.is_server_controller():
            messages.error(request, 'This asset is not a Server Controller.')
            return redirect('freed_systems')
        
        # Get scrapping reason
        reason = request.POST.get('scrapping_reason', '').strip()
        
        # Update asset
        asset.status = 'scrapped'
        asset.scrapped_date = timezone.now()
        asset.scrapping_reason = reason or 'Scrapped from Server Controllers'
        asset.save()
        
        # Log history
        from .services.asset_service import AssetService
        AssetService.log_history(
            asset=asset,
            action='scrapped',
            user=request.user,
            details=f'Server Controller scrapped: {reason or "No reason provided"}'
        )
        
        messages.success(
            request,
            f'Server Controller "{asset.asset_tag}" has been scrapped.'
        )
        return redirect('server_controller_list' + '?tab=freed')


def server_controller_export(request):
    """Export Server Controllers to Excel."""
    # Get server controllers
    queryset = Asset.objects.filter(
        Q(asset_ip_addresses__ip_address__address__startswith='192.168.50.') |
        Q(manual_ip__startswith='192.168.50.'),
        status='active'
    ).distinct().select_related(
        'operating_system', 'team', 'branch'
    ).prefetch_related('asset_ip_addresses__ip_address', 'hosted_vms').order_by('asset_tag')
    
    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Server Controllers'
    
    # Header style
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    
    # Headers
    headers = [
        'Asset Tag', 'System Type', 'IP Address', 'Operating System',
        'Assigned To', 'Team', 'Branch', 'Hardware Serial',
        'Warranty', 'Hyper-V VMs', 'VM IPs'
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Data rows
    for row_idx, asset in enumerate(queryset, 2):
        # Get Hyper-V VMs
        vms = asset.hosted_vms.all()
        vm_names = ', '.join([vm.name for vm in vms]) if vms.exists() else '—'
        vm_ips = ', '.join([vm.get_display_ip() for vm in vms if vm.get_display_ip()]) if vms.exists() else '—'
        
        ws.cell(row=row_idx, column=1, value=asset.asset_tag)
        ws.cell(row=row_idx, column=2, value=asset.system_type)
        ws.cell(row=row_idx, column=3, value=asset.get_ip_address_string())
        ws.cell(row=row_idx, column=4, value=asset.operating_system.name if asset.operating_system else '—')
        ws.cell(row=row_idx, column=5, value=asset.assigned_to or '—')
        ws.cell(row=row_idx, column=6, value=asset.team.name if asset.team else '—')
        ws.cell(row=row_idx, column=7, value=asset.branch.name if asset.branch else '—')
        ws.cell(row=row_idx, column=8, value=asset.hardware_serial_number or '—')
        ws.cell(row=row_idx, column=9, value=asset.warranty_expiration.strftime('%Y-%m-%d') if asset.warranty_expiration else '—')
        ws.cell(row=row_idx, column=10, value=vm_names)
        ws.cell(row=row_idx, column=11, value=vm_ips)
    
    # Adjust column widths
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
    
    # Save to BytesIO
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    # Create response
    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=server_controllers_{timezone.now().strftime("%Y%m%d")}.xlsx'
    
    return response
